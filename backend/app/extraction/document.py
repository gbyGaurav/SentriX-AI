"""General document extraction for DOCX and text formats."""

import io
import re
import logging
from typing import Dict, Any

logger = logging.getLogger(__name__)

URL_REGEX = r'https?://[^\s<>"\']+|www\.[^\s<>"\']+\.[a-zA-Z]{2,}'


def extract_docx_data(docx_bytes: bytes) -> Dict[str, Any]:
    """Extract text, URLs, and metadata from docx file."""
    results = {
        "text": "",
        "urls": [],
        "metadata": {},
        "paragraph_count": 0,
    }

    try:
        from docx import Document

        doc = Document(io.BytesIO(docx_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        table_texts.append(cell.text.strip())

        full_text = "\n".join(paragraphs + table_texts)
        results["text"] = full_text
        results["paragraph_count"] = len(paragraphs)
        
        # Core properties
        cp = doc.core_properties
        results["metadata"] = {
            "author": cp.author or "",
            "title": cp.title or "",
            "created": str(cp.created) if cp.created else "",
            "modified": str(cp.modified) if cp.modified else "",
            "last_modified_by": cp.last_modified_by or "",
        }

        # URLs
        extracted_urls = set(re.findall(URL_REGEX, full_text))
        results["urls"] = list(extracted_urls)

    except Exception as e:
        logger.error(f"Error parsing DOCX: {e}", exc_info=True)
        results["error"] = str(e)

    return results
